"""CSV and Excel reporting: one row per document, a per-case roll-up, an exceptions sheet, a run summary."""
from __future__ import annotations

import csv
import datetime as dt
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

COLUMNS = [
    ("filename", "File"), ("court", "Court"), ("case_number", "Case number"), ("doc_number", "Doc. no."),
    ("filed", "Filed"), ("doc_type", "Document type"), ("pages", "Pages"), ("stamp_page", "Stamp page"),
    ("stamp_total", "Stamp of"), ("is_complete", "Complete"), ("text_source", "Text source"),
    ("status", "Status"), ("review_reasons", "Review reasons"), ("sha256", "SHA-256"), ("new_path", "Organised path"),
]
INK, PAPER = "1B2A41", "F4F1EA"


def _cell(rec: dict, key: str):
    value = rec.get(key)
    if isinstance(value, list):
        return "; ".join(value)
    if isinstance(value, bool):
        return "yes" if value else "no"
    if key == "filed" and value:
        return dt.date.fromisoformat(value)
    return value


def write_csv(records: list[dict], path: str) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([label for _, label in COLUMNS])
        for rec in records:
            writer.writerow(["" if _cell(rec, k) is None else _cell(rec, k) for k, _ in COLUMNS])


def by_case(records: list[dict]) -> list[dict]:
    groups: dict[tuple, list[dict]] = defaultdict(list)
    for rec in records:
        groups[(rec.get("court") or "unresolved", rec.get("case_number") or "unresolved")].append(rec)
    rows = []
    for (court, case), recs in sorted(groups.items()):
        dates = sorted(r["filed"] for r in recs if r.get("filed"))
        rows.append({"court": court, "case_number": case, "documents": len(recs),
                     "pages": sum(r["pages"] for r in recs), "first_filed": dates[0] if dates else None,
                     "last_filed": dates[-1] if dates else None,
                     "types": ", ".join(sorted({r["doc_type"] for r in recs}))})
    return rows


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _table(ws, name: str, nrows: int, ncols: int) -> None:
    ref = f"A1:{get_column_letter(ncols)}{max(nrows, 2)}"
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)


def write_xlsx(records: list[dict], summary: dict, path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Case index"
    ws.append([label for _, label in COLUMNS])
    for rec in records:
        ws.append([_cell(rec, k) for k, _ in COLUMNS])
    _style_header(ws, len(COLUMNS))
    for row in range(2, len(records) + 2):
        ws.cell(row=row, column=5).number_format = "yyyy-mm-dd"
        link = ws.cell(row=row, column=15)
        link.hyperlink = "case_files/" + str(link.value)
        link.font = Font(color="1F4E9E", underline="single")
    _widths(ws, [40, 10, 28, 16, 12, 18, 7, 10, 9, 10, 12, 9, 46, 20, 60])
    _table(ws, "CaseIndex", len(records) + 1, len(COLUMNS))

    cases = by_case(records)
    ws2 = wb.create_sheet("By case")
    ws2.append(["Court", "Case number", "Documents", "Pages", "First filed", "Last filed", "Document types"])
    for row in cases:
        ws2.append([row["court"], row["case_number"], row["documents"], row["pages"],
                    dt.date.fromisoformat(row["first_filed"]) if row["first_filed"] else None,
                    dt.date.fromisoformat(row["last_filed"]) if row["last_filed"] else None, row["types"]])
    _style_header(ws2, 7)
    for r in range(2, len(cases) + 2):
        ws2.cell(row=r, column=5).number_format = "yyyy-mm-dd"
        ws2.cell(row=r, column=6).number_format = "yyyy-mm-dd"
    _widths(ws2, [12, 30, 11, 8, 12, 12, 40])
    _table(ws2, "ByCase", len(cases) + 1, 7)

    ws3 = wb.create_sheet("Exceptions")
    ws3.append(["File", "Case number", "Document type", "Review reasons"])
    exceptions = [r for r in records if r["status"] != "ok"]
    for rec in exceptions:
        ws3.append([rec["filename"], rec.get("case_number"), rec["doc_type"], "; ".join(rec["review_reasons"])])
    _style_header(ws3, 4)
    _widths(ws3, [42, 28, 20, 80])
    if exceptions:
        _table(ws3, "Exceptions", len(exceptions) + 1, 4)

    ws4 = wb.create_sheet("Run summary")
    ws4.append(["Measure", "Value"])
    for key, value in summary.items():
        ws4.append([key, value])
    _style_header(ws4, 2)
    _widths(ws4, [44, 40])
    wb.save(path)
