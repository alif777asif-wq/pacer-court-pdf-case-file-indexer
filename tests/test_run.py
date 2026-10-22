"""Integration tests against the pinned real corpus."""
import csv
import hashlib
import json
import os

import openpyxl

from courtfile.organize import filename_hints, plan_paths, target_path


def test_every_pdf_is_indexed_and_reconciles(run_output, corpus_dir):
    results, records, out = run_output
    assert results["verification"]["reconciliation"]["consistent"] is True
    assert len(records) == len(list(corpus_dir.glob("*.pdf")))


def test_corpus_matches_lock_file_checksums(corpus_dir):
    lock = json.loads((corpus_dir.parents[1] / "corpus.lock.json").read_text())
    for entry in lock["files"]:
        assert hashlib.sha256((corpus_dir / entry["local_name"]).read_bytes()).hexdigest() == entry["sha256"]


def test_stamp_fields_agree_with_independent_filename_source(run_output):
    check = run_output[0]["verification"]["filename_vs_text"]
    assert check["fields_compared"] >= 10
    assert check["mismatches"] == []


def test_case_numbers_agree_across_two_extractors(run_output):
    results, records, _ = run_output
    bad = [d for d in results["verification"]["cross_extractor"]["disagreements"]
           if d.get("poppler", [0])[0] != d.get("pdfminer", [1])[0]]
    assert bad == []


def test_overprinted_stamp_is_routed_to_review_not_guessed(run_output):
    _, records, _ = run_output
    rec = next(r for r in records if "cand.203070" in r["filename"])
    assert rec["doc_number_text"] is None and rec["status"] == "review"
    assert rec["doc_number"] == "27"            # recovered from the file name and labelled as such
    assert rec["doc_number_source"] == "filename"


def test_second_extractor_value_for_overprinted_stamp_is_wrong(run_output):
    """The check is meaningful: pdfminer reads the overprint as document 26, the file name says 27."""
    dis = {d["file"]: d for d in run_output[0]["verification"]["cross_extractor"]["disagreements"]}
    assert dis["doctor__gov.uscourts.cand.203070.27.0.pdf"]["pdfminer"][1] == "26"


def test_scanned_page_is_recovered_by_ocr(run_output):
    _, records, _ = run_output
    rec = next(r for r in records if r["filename"] == "xray__no_bad_redactions.4.1.pdf")
    assert rec["ocr_pages"] == [1] and rec["text_chars"] > 300 and rec["case_number"] == "2:15-cv-05642-CAS-JC"


def test_classifier_reference_accuracy_and_known_miss(run_output):
    ref = run_output[0]["verification"]["reference_labels"]
    assert ref["documents_scored"] == 13 and ref["correct"] >= 12
    assert [m["file"] for m in ref["misses"]] == ["xray__no_bad_redactions.4.1.pdf"]


def test_planted_errors_are_caught(run_output):
    st = run_output[0]["verification"]["planted_error_selftest"]
    assert st["errors_planted"] == 2 and st["errors_caught"] == 2


def test_organised_copies_are_byte_identical(run_output):
    assert run_output[0]["verification"]["copy_integrity"]["mismatches"] == []


def test_every_personal_data_flag_has_a_documented_manual_decision(run_output):
    screen = run_output[0]["personal_data_screen"]
    assert screen["flagged_unreviewed"] == []


def test_organised_paths_are_portable_and_unique(run_output):
    _, records, _ = run_output
    paths = [r["new_path"] for r in records]
    assert len(set(paths)) == len(paths)
    assert all(":" not in p and " " not in p for p in paths)


def test_excel_workbook_round_trip(run_output):
    _, records, out = run_output
    wb = openpyxl.load_workbook(out / "case_index.xlsx")
    assert wb.sheetnames == ["Case index", "By case", "Exceptions", "Run summary"]
    ws = wb["Case index"]
    assert ws.max_row - 1 == len(records)
    for row in ws.iter_rows(min_row=2):
        assert os.path.exists(out / "case_files" / row[14].value)
    with open(out / "case_index.csv", newline="", encoding="utf-8") as fh:
        assert sum(1 for _ in csv.reader(fh)) - 1 == len(records)


def test_path_collisions_get_a_checksum_suffix():
    base = {"court": "cod", "case_number": "1:21-cv-01061-KLM", "doc_number": "1", "doc_type": "Excerpt",
            "filed": "2021-04-16", "stamp_page": None, "stamp_total": None, "is_complete": None}
    a = dict(base, sha256="a" * 64)
    b = dict(base, sha256="b" * 64)
    plan_paths([a, b])
    assert a["new_path"] != b["new_path"] and target_path(a).count(":") == 0


def test_filename_hints():
    assert filename_hints("doctor__ca10_010110462922.pdf") == {"court": "ca10", "doc_number": "010110462922"}
    assert filename_hints("gov.uscourts.cacd.652774.40.0.pdf") == {"court": "cacd", "doc_number": "40"}
    assert filename_hints("random-name.pdf") == {}
